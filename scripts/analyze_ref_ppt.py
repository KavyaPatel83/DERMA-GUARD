import sys
import io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

from pptx import Presentation
from pptx.util import Pt

prs = Presentation(r'd:\VIT BOOKS\PROJECT 1\Healthsphere_1\Project-1__Review2_Final.pptx')
print(f'Slides: {len(prs.slides)}')

slides_list = list(prs.slides)
for i, slide in enumerate(slides_list):
    print(f'\n=== SLIDE {i+1} ===')
    bg = slide.background
    fill = bg.fill
    print(f'  BG fill type: {fill.type}')
    try:
        print(f'  BG fore_color rgb: #{fill.fore_color.rgb}')
    except: pass

    for j, shape in enumerate(slide.shapes):
        try:
            sfill = shape.fill
            try:
                out = f'  Shape[{j}] name="{shape.name}" fill_type={sfill.type}'
                try: out += f' color=#{sfill.fore_color.rgb}'
                except: pass
                print(out)
            except: pass
        except: pass

        if shape.has_text_frame:
            for k, para in enumerate(shape.text_frame.paragraphs[:2]):
                txt = para.text.strip().encode('ascii','replace').decode()
                if not txt: continue
                for run in para.runs[:1]:
                    try:
                        rgb = str(run.font.color.rgb)
                    except:
                        rgb = 'theme/inherited'
                    try:
                        sz = run.font.size.pt if run.font.size else 'inh'
                    except:
                        sz = 'inh'
                    bold = getattr(run.font, 'bold', None)
                    fname = getattr(run.font, 'name', 'inh')
                    print(f'    Text: "{txt[:60]}" | font={fname} sz={sz} bold={bold} color={rgb}')
                    break
